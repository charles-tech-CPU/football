#!/usr/bin/env python3
"""
Importe les lignes de "Calendrier championnat" a date connue mais adversaire(s)
pas encore tire au sort (jetons de bracket type HF/QF/3E/W1, ou repetition du
meme nom reel dans les deux colonnes club) : ignorees par import_excel.py car
non exploitables telles quelles. On les importe quand meme avec une equipe
"A DETERMINER" generique a la place du/des cote(s) inconnu(s), pour que
l'utilisateur puisse remplacer par la bonne equipe des que le tirage est fait.

Usage : python3 extract_pending_draws.py <fichier.xlsx> <competitions.json> <teams.json> <out.sql>
"""
import json
import re
import sys
import unicodedata
from datetime import datetime

import openpyxl

SEASON = 2027
PLACEHOLDER_RE = re.compile(r"^(HF|QF|DF|F|\d+E|W\d+|L\d+)$", re.IGNORECASE)
# 2 equipes generiques distinctes (pas une seule) : la base interdit qu'un
# match oppose une equipe a elle-meme (chk_match_teams_distinct), ce qui
# arriverait des qu'aucun des deux cotes n'est encore connu (ex: "HF" vs "HF").
PLACEHOLDER_TEAM_1 = "A DETERMINER (1)"
PLACEHOLDER_TEAM_2 = "A DETERMINER (2)"

CLUB_NOISE_RE = re.compile(
    r"\b(FC|CF|AC|AS|SC|SK|FK|NK|BK|CD|CS|UD|RC|CA|OGC|OL|PSG|KF|HNK|MSK|MFK|CP|"
    r"CLUB|FOOTBALL|CLUB DE FUTBOL|CALCIO|ATLETICO|ATHLETIC)\b",
    re.IGNORECASE,
)


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def normalize(name):
    n = strip_accents(name).upper()
    n = n.replace(".", " ").replace("-", " ")
    n = CLUB_NOISE_RE.sub(" ", n)
    n = re.sub(r"[^A-Z0-9 ]", " ", n)
    n = re.sub(r"\s+", " ", n).strip()
    return n


def sql_escape(v):
    return str(v).replace("'", "''")


NAME_ALIASES = {
    "TROMSO": "TROMSO IL",
}


def build_team_resolver(teams):
    by_norm = {}
    for t in teams:
        by_norm.setdefault(normalize(t["name"]), []).append(t["name"])

    def resolve(name):
        if name in NAME_ALIASES:
            return NAME_ALIASES[name]
        exact = [t["name"] for t in teams if t["name"] == name]
        if exact:
            return exact[0]
        matches = by_norm.get(normalize(name), [])
        if len(matches) == 1:
            return matches[0]
        return None

    return resolve


def main():
    if len(sys.argv) != 5:
        print("Usage: extract_pending_draws.py <fichier.xlsx> <competitions.json> <teams.json> <out.sql>", file=sys.stderr)
        sys.exit(1)

    xlsx_path, comps_path, teams_path, out_path = sys.argv[1:5]

    with open(comps_path, encoding="utf-8") as f:
        competitions = json.load(f)
    with open(teams_path, encoding="utf-8") as f:
        teams = json.load(f)

    resolve_team = build_team_resolver(teams)

    existing_leagues = {(c["country"] or "").strip().upper(): c for c in competitions if c["type"] == "LEAGUE"}
    existing_cups = {(c["country"] or "").strip().upper(): c for c in competitions if c["type"] == "DOMESTIC_CUP"}

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["Calendrier championnat"]

    new_cup_competitions = {}  # country_key -> display_country
    rows_out = []
    unresolved_names = set()

    for row in range(2, ws.max_row + 1):
        pays = ws.cell(row=row, column=1).value
        journee = ws.cell(row=row, column=2).value
        date = ws.cell(row=row, column=3).value
        horaire = ws.cell(row=row, column=4).value
        club1 = ws.cell(row=row, column=5).value
        club2 = ws.cell(row=row, column=6).value

        if not pays or not club1 or not club2 or not isinstance(date, datetime):
            continue

        club1, club2 = str(club1).strip(), str(club2).strip()
        tok1 = bool(PLACEHOLDER_RE.match(club1))
        tok2 = bool(PLACEHOLDER_RE.match(club2))
        self_repeat = club1 == club2 and not tok1

        if not (tok1 or tok2 or self_repeat):
            continue  # ligne normale, deja geree par import_excel.py

        country_key = str(pays).strip().upper()
        display_country = str(pays).strip().title()

        # -- resolution des 2 cotes --
        def side(name, is_token):
            if is_token:
                return None, name  # None = placeholder
            resolved = resolve_team(name)
            if resolved is None:
                unresolved_names.add(name)
                return None, name
            return resolved, name

        team1_name, raw1 = side(club1, tok1)
        team2_name, raw2 = (None, club2) if self_repeat else side(club2, tok2)

        # -- competition + round_label --
        if isinstance(journee, (int, float)):
            comp = existing_leagues.get(country_key)
            round_label = f"J{int(journee)}"
            comp_kind = "LEAGUE"
        elif journee == "CUP":
            comp = existing_cups.get(country_key)
            round_label = f"Coupe nationale (a determiner : {raw1}-{raw2})"
            comp_kind = "CUP"
        elif journee == "BARRAGE":
            comp = existing_leagues.get(country_key)
            round_label = f"Barrage (a determiner : {raw1}-{raw2})"
            comp_kind = "LEAGUE"
        elif journee == "B EUR":
            comp = existing_leagues.get(country_key)
            round_label = f"Barrage Europe (a determiner : {raw1}-{raw2})"
            comp_kind = "LEAGUE"
        else:
            continue

        if comp is None and comp_kind == "CUP":
            new_cup_competitions[country_key] = display_country
            comp_code = f"{country_key}-CUP"
        elif comp is None:
            # championnat introuvable (pays sans championnat importe) : on saute, ne devrait pas arriver
            print(f"ATTENTION : pas de championnat pour {pays} (ligne {row}), ignoree.", file=sys.stderr)
            continue
        else:
            comp_code = comp["code"]

        rows_out.append({
            "comp_code": comp_code,
            "round_label": round_label,
            "date": date.date().isoformat(),
            "time": None,
            "team1": team1_name,
            "team2": team2_name,
        })

    print(f"{len(rows_out)} lignes a importer, {len(new_cup_competitions)} nouvelles competitions de coupe, "
          f"{len(unresolved_names)} noms non resolus", file=sys.stderr)
    for n in sorted(unresolved_names):
        print("  NON RESOLU:", n, file=sys.stderr)

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("-- Matchs a date connue mais adversaire(s) pas encore tire au sort (coupes nationales,\n")
        f.write("-- barrages promotion/relegation, barrages qualification europeenne) : importes avec\n")
        f.write("-- une equipe 'A DETERMINER' a la place du/des cote(s) inconnu(s), a remplacer par la\n")
        f.write("-- bonne equipe une fois le tirage/la position finale connue (voir import/extract_pending_draws.py).\n\n")

        f.write(f"INSERT INTO team (name, country) VALUES ('{PLACEHOLDER_TEAM_1}', NULL);\n")
        f.write(f"INSERT INTO team (name, country) VALUES ('{PLACEHOLDER_TEAM_2}', NULL);\n\n")

        for country_key, display_country in sorted(new_cup_competitions.items()):
            comp_code = f"{country_key}-CUP"
            name = f"{display_country} - Coupe nationale {SEASON}"
            f.write(
                f"INSERT INTO competition (code, name, type, country, season) VALUES "
                f"('{sql_escape(comp_code)}', '{sql_escape(name)}', 'DOMESTIC_CUP', '{sql_escape(display_country)}', {SEASON});\n"
            )
        f.write("\n")

        for r in rows_out:
            team1_sub = (
                f"(SELECT id FROM team WHERE name = '{sql_escape(r['team1'])}')"
                if r["team1"] else f"(SELECT id FROM team WHERE name = '{PLACEHOLDER_TEAM_1}')"
            )
            team2_sub = (
                f"(SELECT id FROM team WHERE name = '{sql_escape(r['team2'])}')"
                if r["team2"] else f"(SELECT id FROM team WHERE name = '{PLACEHOLDER_TEAM_2}')"
            )
            f.write(
                "INSERT INTO match (competition_id, round_label, date, time, team1_id, team2_id, score1, score2, status) VALUES ("
                f"(SELECT id FROM competition WHERE code = '{sql_escape(r['comp_code'])}' AND season = {SEASON}), "
                f"'{sql_escape(r['round_label'])}', '{r['date']}', NULL, "
                f"{team1_sub}, {team2_sub}, NULL, NULL, 'SCHEDULED');\n"
            )

    print(f"OK -> {out_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
